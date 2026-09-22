from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.ai import stub_analyze
from app.catalogs.product_catalog import (
    build_product_list_attachment,
    classify_category_interests,
    customer_interest_keys,
    load_catalog_yaml,
    merge_customer_interests,
    render_product_list_email,
    validate_product_list_email,
)
from app.catalogs.products import (
    canonical_product_code,
    find_product_codes,
    load_product_aliases,
)
from app.db import Customer, Product, ProductCategory


def test_catalog_yaml_has_audited_categories_and_unique_internal_products() -> None:
    payload = load_catalog_yaml()
    categories = payload["categories"]
    products = payload["products"]

    assert [item["key"] for item in categories] == [
        "industrial_silanes",
        "pharmaceutical",
        "rubber_plastics",
        "acetylacetone_salts",
        "silicone_oil",
    ]
    keys = {item["key"] for item in categories}
    codes = [item["code"] for item in products]
    assert len(codes) == len(set(codes)) == 17
    assert all(item["category"] in keys for item in products)
    assert any(item["code"] == "DEMO-P100" for item in products)
    assert any(item["code"] == "DEMO-P302" for item in products)
    assert all(item.get("cas_no") is None for item in products)
    assert any(item["code"] == "DEMO-P800" for item in products)
    assert any(item["code"] == "DEMO-GRADE-A(99%)" for item in products)
    hidden = [item for item in products if item.get("catalog_visible") is False]
    assert [item["code"] for item in hidden] == ["DEMO-GRADE-A(99%)"]
    public_codes = [
        item.get("catalog_code", item["code"])
        for item in products
        if item.get("catalog_visible", True)
    ]
    assert len(public_codes) == len(set(public_codes)) == 16
    assert "DEMO-GRADE-A" in public_codes
    assert all(item.get("content") != "-" for item in products)


def test_unreviewed_database_products_default_to_catalog_hidden() -> None:
    default = Product.__table__.c.catalog_visible.default
    assert default is not None and default.arg is False


def test_category_keyword_classification() -> None:
    assert classify_category_interests("工业硅烷") == ["industrial_silanes"]
    assert classify_category_interests("we are interested in industrial silane") == [
        "industrial_silanes"
    ]
    assert classify_category_interests("医药 API") == ["pharmaceutical"]
    assert classify_category_interests("橡塑 PVC heat stabilizers antioxidants") == [
        "rubber_plastics"
    ]
    assert classify_category_interests("rubber and plastics") == ["rubber_plastics"]
    assert classify_category_interests("hello world") == []
    assert classify_category_interests("工业硅烷 and 医药") == [
        "industrial_silanes",
        "pharmaceutical",
    ]


def test_customer_interest_metadata_merges_and_deduplicates() -> None:
    customer = Customer(company_name="Example Buyer", metadata_json={})
    assert customer_interest_keys(customer) == []

    merge_customer_interests(
        customer,
        [
            {
                "category_key": "industrial_silanes",
                "category_name": "Industrial Silanes",
                "source": "full_customer_workbook",
                "value": "工业硅烷",
            }
        ],
    )
    merge_customer_interests(
        customer,
        [
            {
                "category_key": "industrial_silanes",
                "category_name": "Industrial Silanes",
                "source": "full_customer_workbook",
                "value": "工业硅烷",
            }
        ],
    )
    assert customer_interest_keys(customer) == ["industrial_silanes"]
    assert len(customer.metadata_json["interests"]) == 1


def test_new_aliases_resolve_product_codes() -> None:
    load_product_aliases.cache_clear()
    try:
        assert canonical_product_code("DEMO P100") == "DEMO-P100"
        assert canonical_product_code("P100") == "DEMO-P100"
        assert canonical_product_code("BCP") == "DEMO-BCP"
        assert canonical_product_code("P302") == "DEMO-P302"
        assert canonical_product_code("Polymer Additive 01") == "DEMO-AO-01"
        assert canonical_product_code("DEMO-GRADE-A") == "DEMO-GRADE-A(98%)"
        assert find_product_codes("Please quote Reagent B for our plant.") == ["DEMO-BCP"]
    finally:
        load_product_aliases.cache_clear()


def test_product_list_email_rendering_is_deterministic_and_price_free() -> None:
    category = ProductCategory(key="industrial_silanes", name="Industrial Silanes")
    products = [
        Product(
            code="DEMO-P100",
            catalog_code="DEMO-P100",
            name="Industrial Coupling Agent 100",
            brand="Demo",
            cas_no=None,
            content="Demo grade",
            series="Demonstration Series",
            sort_order=1,
            id=1,
        ),
        Product(
            code="DEMO-P500",
            catalog_code="DEMO-P500",
            name="Industrial Treatment Agent 500",
            brand="Demo",
            cas_no=None,
            content="Demo grade",
            series="Demonstration Series",
            sort_order=2,
            id=2,
        ),
        Product(
            code="DEMO-GRADE-A(99%)",
            catalog_code="DEMO-GRADE-A",
            catalog_visible=False,
            name="Demonstration Grade A",
            cas_no=None,
            content="99%",
            sort_order=3,
            id=3,
        ),
    ]

    text, html_body = render_product_list_email(
        contact_name="Alice Buyer",
        category=category,
        products=products,
        subject="Re: product list",
        signature_text="Best regards,\nExample Sales Team",
        signature_html="<p>Best regards,</p><p>Example Sales Team</p>",
    )

    assert "Dear Alice Buyer," in text
    assert "Industrial Silanes" in text
    assert "DEMO-P100" in text
    assert "Demonstration Series" in text
    assert "US$" not in text and "USD" not in text
    assert "<table" in html_body
    assert "DEMO-P500" in html_body
    assert "DEMO-GRADE-A" not in text and "DEMO-GRADE-A" not in html_body
    assert text.count("Best regards,") == 1
    assert html_body.count("Best regards,") == 1


def test_product_list_email_rejects_money_and_commitments() -> None:
    with pytest.raises(ValueError, match="monetary"):
        validate_product_list_email("The price is USD 12.50 per kg.")
    with pytest.raises(ValueError, match="commitment"):
        validate_product_list_email("We guarantee delivery.")


def test_product_list_workbook_uses_only_curated_values_and_keeps_missing_cas_blank() -> None:
    category = ProductCategory(key="pharmaceutical", name="Pharmaceuticals")
    products = [
        Product(
            code="DEMO-P400",
            catalog_code="DEMO-P400",
            name="General Reagent 400",
            cas_no=None,
            content="99%",
            series="Demonstration Series",
            sort_order=1,
            id=1,
        ),
        Product(
            code="DEMO-P500",
            catalog_code="DEMO-P500",
            name="Industrial Treatment Agent 500",
            cas_no=None,
            content="Demo grade",
            series="Demonstration Series",
            sort_order=2,
            id=2,
        ),
        Product(
            code="DEMO-GRADE-A(99%)",
            catalog_code="DEMO-GRADE-A",
            catalog_visible=False,
            name="Demonstration Grade A",
            cas_no=None,
            content="99%",
            sort_order=3,
            id=3,
        ),
    ]

    attachment = build_product_list_attachment(
        category=category,
        products=products,
        file_format="xlsx",
    )
    workbook = load_workbook(BytesIO(attachment.payload), data_only=False)
    sheet = workbook["Product List"]

    assert attachment.filename == "Example_Chemicals_pharmaceutical_product_list.xlsx"
    assert sheet["A1"].value == "No."
    assert sheet["E1"].value == "CAS No."
    assert sheet["C2"].value == "DEMO-P400"
    assert sheet["E2"].value is None
    assert sheet["C3"].value == "DEMO-P500"
    assert sheet["E3"].value is None
    assert sheet.max_row == 3
    assert all(cell.value != "'-" for row in sheet.iter_rows() for cell in row)


def test_attached_product_list_email_requires_a_real_attachment_context() -> None:
    category = ProductCategory(key="pharmaceutical", name="Pharmaceuticals")
    product = Product(
        code="DEMO-P400",
        catalog_code="DEMO-P400",
        name="General Reagent 400",
        cas_no=None,
        content="99%",
        sort_order=1,
        id=1,
    )

    text, _ = render_product_list_email(
        contact_name="Alice",
        category=category,
        products=[product],
        subject="Product list in Excel",
        signature_text="Best regards,\nExample Sales Team",
        signature_html="<p>Best regards,</p><p>Example Sales Team</p>",
        attachment_filename="Example_Chemicals_pharmaceutical_product_list.xlsx",
    )

    assert "Please find attached" in text
    with pytest.raises(ValueError, match="attachment claim"):
        validate_product_list_email("Please find attached our product list.")


def test_stub_analysis_classifies_category_only_inquiries_as_product_list() -> None:
    analysis = stub_analyze(
        "Re: inquiry",
        "We are interested in industrial silane. Please send your product list.",
        [],
    )
    assert analysis.intent.value == "product_list_request"
    assert analysis.product_code is None
    assert analysis.intent_confidence == 0.95
    assert analysis.product_confidence == 0.93
    assert "product_code" not in analysis.missing_fields

    analysis = stub_analyze(
        "Re: inquiry",
        "Please send your full product catalog.",
        [],
    )
    assert analysis.intent.value == "product_list_request"

    analysis = stub_analyze(
        "Re: inquiry",
        "PRODUCT DEMO-P100. Please quote 100 kg.",
        [],
    )
    assert analysis.intent.value == "quote_request"
    assert analysis.product_code == "DEMO-P100"


def test_explicit_product_code_extraction_ignores_list_catalog_words() -> None:
    from app.services import _explicit_product_codes

    assert _explicit_product_codes("Please send your product list for industrial silane.") == []
    assert _explicit_product_codes("Please send your full product catalog.") == []
    assert _explicit_product_codes("Please send the product brochure.") == []
    assert _explicit_product_codes("Please send a sample of your industrial silane products.") == []
    assert _explicit_product_codes("PRODUCTS for the rubber industry, please.") == []
    assert _explicit_product_codes("PRODUCT DEMO-P100. Please send your product list.") == [
        "DEMO-P100"
    ]
    assert _explicit_product_codes("Please quote PRODUCT WIDGET-100 quantity 100 kg.") == [
        "WIDGET-100"
    ]


def test_stub_analysis_extracts_multi_product_requests_with_quantities() -> None:
    analysis = stub_analyze(
        "Re: quotation",
        "Please quote PRODUCT WIDGET-100 100 kg and PRODUCT WIDGET-200 200 kg.",
        [],
    )
    assert len(analysis.product_requests) == 2
    by_code = {line.product_code: line.quantity for line in analysis.product_requests}
    assert by_code.get("WIDGET-100") == 100
    assert by_code.get("WIDGET-200") == 200

    partial = stub_analyze(
        "Re: quotation",
        "Please quote PRODUCT WIDGET-100 100 kg and PRODUCT WIDGET-200.",
        [],
    )
    by_code = {line.product_code: line.quantity for line in partial.product_requests}
    assert by_code.get("WIDGET-100") == 100
    assert by_code.get("WIDGET-200") is None
